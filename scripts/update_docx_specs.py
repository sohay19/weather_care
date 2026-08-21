from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"


def set_cell_border(cell, *, color: str = "AFC4DE", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_width(cell, width_cm: float) -> None:
    width = Cm(width_cm)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def apply_document_accessibility(document: Document, title: str) -> None:
    document.core_properties.title = title
    for table in document.tables:
        if table.rows:
            set_repeat_table_header(table.rows[0])


def style_environment_table(document: Document) -> None:
    table = next(
        table
        for table in document.tables
        if len(table.rows) >= 2
        and len(table.columns) == 4
        and table.cell(0, 1).text.strip() == "원본 시간"
    )
    table.autofit = False
    widths = (2.5, 4.4, 5.3, 4.8)
    set_repeat_table_header(table.rows[0])

    for row_index, row in enumerate(table.rows):
        row.height = None
        for column_index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[column_index])
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                shading = cell._tc.get_or_add_tcPr().first_child_found_in("w:shd")
                if shading is None:
                    shading = OxmlElement("w:shd")
                    cell._tc.get_or_add_tcPr().append(shading)
                shading.set(qn("w:fill"), "D9E8F5")

            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.paragraph_format.line_spacing = 1.08
                for run in paragraph.runs:
                    run.font.size = Pt(9.5 if row_index == 0 else 9)
                    if row_index == 0:
                        run.bold = True


def update_cover_date(document: Document) -> None:
    for table in document.tables:
        for row in table.rows:
            cells = row.cells
            for index, cell in enumerate(cells[:-1]):
                if cell.text.strip() == "기준일":
                    cells[index + 1].text = "2026-08-21"
                    return
    raise RuntimeError("기준일 셀을 찾지 못했습니다.")


def remove_hard_break_before(document: Document, heading_text: str) -> None:
    heading = next(p for p in document.paragraphs if p.text.strip() == heading_text)
    previous = heading._p.getprevious()
    if previous is not None and previous.tag == qn("w:p"):
        has_page_break = bool(previous.xpath('.//w:br[@w:type="page"]'))
        if has_page_break:
            previous.getparent().remove(previous)


def set_page_break_before(document: Document, text: str, value: bool) -> None:
    paragraph = next(p for p in document.paragraphs if p.text.strip() == text)
    paragraph.paragraph_format.page_break_before = value


def update_app_spec() -> None:
    path = DOCS / "날씨챙겨_앱_개발명세.docx"
    document = Document(path)
    old = (
        "Main 새로고침은 상단 아이콘 버튼으로 제공한다. "
        "Pull-to-refresh를 위해 스크롤을 억지로 추가하지 않는다."
    )
    new = (
        "Main을 포함한 모든 1차 탭 화면은 화면을 아래로 당기는 RefreshIndicator로 "
        "새로고침한다. 별도 새로고침 아이콘 버튼은 두지 않는다."
    )
    paragraph = next(
        p for p in document.paragraphs if p.text.strip() in {old, new}
    )
    if paragraph.text.strip() != new:
        paragraph.text = new
        paragraph.style = "List Bullet"
    apply_document_accessibility(document, "날씨챙겨 앱 개발명세")
    document.save(path)


def update_server_spec() -> None:
    path = DOCS / "날씨챙겨_서버_개발명세.docx"
    document = Document(path)
    paragraph = next(p for p in document.paragraphs if p.text.strip() == "PARASOL / 양산")
    paragraph.paragraph_format.space_after = Pt(7)
    apply_document_accessibility(document, "날씨챙겨 서버 개발명세")
    document.save(path)


def update_catalog_v10() -> None:
    path = DOCS / "날씨챙겨_시간대별_판단기준_및_문구카탈로그_v1.0.docx"
    document = Document(path)
    update_cover_date(document)
    for heading in (
        "7. Recommendation 문구 카탈로그 — 계속",
        "8. Lifestyle 문구 카탈로그 — 계속",
        "10. D1 이전 — seed와 조회",
        "11. 검증 기준과 구현 순서",
    ):
        remove_hard_break_before(document, heading)
    set_page_break_before(document, "10. D1 이전 — seed와 조회", True)
    style_environment_table(document)
    apply_document_accessibility(
        document,
        "날씨챙겨 시간대별 판단 기준 및 문구 카탈로그 v1.0",
    )
    document.save(path)


def update_catalog_v11() -> None:
    path = DOCS / "날씨챙겨_시간대별_판단기준_및_문구카탈로그_v1.1.docx"
    document = Document(path)
    update_cover_date(document)
    for heading in (
        "6. 세탁·환기 기준",
        "12. 파생 데이터와 시간 창 계산",
        "14. v1.1 추가 문구 카탈로그",
    ):
        set_page_break_before(document, heading, False)
    style_environment_table(document)
    apply_document_accessibility(
        document,
        "날씨챙겨 시간대별 판단 기준 및 문구 카탈로그 v1.1",
    )
    document.save(path)


def main() -> None:
    update_app_spec()
    update_server_spec()
    update_catalog_v10()
    update_catalog_v11()


if __name__ == "__main__":
    main()
