from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "스토어_등록정보.md"
OUTPUT = ROOT / "날씨챙겨_스토어_등록정보.docx"
ICON = ROOT / "appIcon.png"

FONT = "Malgun Gothic"
BLUE = "4C77A4"
PALE_BLUE = "EDF4F8"
LIGHT_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(92, 104, 116)


def set_run_font(run, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clean_inline(text: str) -> str:
    text = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1 (\2)", text)
    text = text.replace("**", "").replace("`", "")
    return text


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor(29, 40, 50)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.16

    title = doc.styles["Title"]
    title.font.name = FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_before = Pt(5)
    title.paragraph_format.space_after = Pt(10)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = FONT
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading1.font.size = Pt(16)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(7)
    heading1.paragraph_format.keep_with_next = True

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = FONT
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    heading2.font.size = Pt(12.5)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2.paragraph_format.space_before = Pt(11)
    heading2.paragraph_format.space_after = Pt(5)
    heading2.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(10.3)
        style.paragraph_format.space_after = Pt(3.5)
        style.paragraph_format.line_spacing = 1.13


def set_cell_padding(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LIGHT_GRAY)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    column_count = len(rows[0])
    table = doc.add_table(rows=1, cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    usable_width = 6.98
    width_ratios = {
        2: (0.31, 0.69),
        3: (0.22, 0.36, 0.42),
        4: (0.20, 0.36, 0.23, 0.21),
        5: (0.16, 0.27, 0.21, 0.18, 0.18),
    }.get(column_count, tuple(1 / column_count for _ in range(column_count)))

    for col_index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[col_index]
        cell.width = Inches(usable_width * width_ratios[col_index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_padding(cell)
        shade_cell(cell, BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(clean_inline(value))
        set_run_font(run, size=9.2, bold=True)
        run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.width = Inches(usable_width * width_ratios[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_padding(cell)
            if row_index % 2 == 0:
                shade_cell(cell, PALE_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if column_count <= 3 and col_index == 0 and len(value) < 24:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(clean_inline(value).replace("<br>", "\n"))
            set_run_font(run, size=9.1, bold=False)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run("날씨챙겨 스토어 등록정보  ")
    set_run_font(run, size=8.5)
    run.font.color.rgb = TEXT_GRAY
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, value, end))


def add_plain_paragraph(doc: Document, text: str, *, code_block: bool = False) -> None:
    paragraph = doc.add_paragraph()
    if code_block:
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.right_indent = Inches(0.05)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.12
    run = paragraph.add_run(clean_inline(text))
    set_run_font(run, size=10.3)
    if code_block:
        run.font.color.rgb = RGBColor(41, 56, 68)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        rows.append(row)
        index += 1
    if len(rows) > 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        rows.pop(1)
    return rows, index


def build_document() -> None:
    doc = Document()
    style_document(doc)
    add_page_number(doc.sections[0])

    if ICON.exists():
        icon_paragraph = doc.add_paragraph()
        icon_paragraph.paragraph_format.space_after = Pt(4)
        icon_paragraph.add_run().add_picture(str(ICON), width=Inches(0.72))

    title = doc.add_paragraph(style="Title")
    title.add_run("날씨챙겨 스토어 등록정보")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run("Google Play 및 App Store 제출용 원고와 체크리스트")
    set_run_font(subtitle_run, size=11.5, bold=False)
    subtitle_run.font.color.rgb = TEXT_GRAY

    lines = SOURCE.read_text(encoding="utf-8").splitlines()[1:]
    index = 0
    in_code = False
    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()

        if stripped.startswith("```"):
            in_code = not in_code
            index += 1
            continue
        if in_code:
            if stripped:
                add_plain_paragraph(doc, raw, code_block=True)
            else:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(2)
            index += 1
            continue
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue
        if stripped.startswith("## "):
            paragraph = doc.add_paragraph(clean_inline(stripped[3:]), style="Heading 1")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if stripped.startswith("### "):
            paragraph = doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2")
            paragraph.paragraph_format.keep_with_next = True
            index += 1
            continue
        if stripped.startswith("- [ ] "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
            paragraph.paragraph_format.space_after = Pt(3.5)
            run = paragraph.add_run("☐ " + clean_inline(stripped[6:]))
            set_run_font(run, size=10.3)
            index += 1
            continue
        if stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(stripped[2:]))
            set_run_font(run, size=10.3)
            index += 1
            continue
        ordered = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if ordered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.28)
            paragraph.paragraph_format.first_line_indent = Inches(-0.24)
            paragraph.paragraph_format.space_after = Pt(3.5)
            run = paragraph.add_run(f"{ordered.group(1)}.  {clean_inline(ordered.group(2))}")
            set_run_font(run, size=10.3)
            index += 1
            continue

        add_plain_paragraph(doc, stripped)
        index += 1

    doc.core_properties.title = "날씨챙겨 스토어 등록정보"
    doc.core_properties.subject = "Google Play 및 App Store 제출용 원고와 체크리스트"
    doc.core_properties.author = "코드소하(CODESOHA)"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
